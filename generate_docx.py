from docx import Document

document = Document()

document.add_heading('Concept Note: EcoSort AI', 0)

document.add_heading('1. Project Title', level=2)
document.add_paragraph('EcoSort AI: Intelligent Waste Segregation Agent')

document.add_heading('2. Target Sustainable Development Goal (SDG)', level=2)
document.add_paragraph('SDG 12: Responsible Consumption and Production')

document.add_heading('3. The Problem', level=2)
document.add_paragraph(
    'Globally, recycling efforts are severely hampered by "wishcycling" and improper waste '
    'segregation at the household level. When consumers place non-recyclable items into '
    'recycling bins, entire batches of recyclables become contaminated and are diverted to '
    'landfills. This problem is exacerbated by confusing packaging labels and municipal '
    'guidelines that vary widely from city to city. Citizens often want to do the right thing '
    'but lack immediate, accessible guidance at the moment of disposal.'
)

document.add_heading('4. The Proposed Solution', level=2)
document.add_paragraph(
    'EcoSort AI is a mobile application powered by Agentic AI designed to eliminate confusion in '
    'household waste disposal. Users simply point their smartphone camera at an item they wish to throw '
    'away. The app instantly identifies the material and provides precise, localized instructions on how '
    'to properly dispose of it based on their municipality\'s specific rules.'
)

document.add_heading('How Agentic AI is Used (Multi-Agent Architecture)', level=3)
document.add_paragraph('Our solution relies on a multi-agent AI system coordinating seamlessly in the background:')
document.add_paragraph('1. Vision Agent: Processes the camera input using computer vision models to accurately identify the object (e.g., "greasy pizza box", "PET plastic bottle").', style='List Number')
document.add_paragraph('2. Localization Researcher Agent: Autonomously queries a vector database containing the specific waste management rules for the user\'s zip code/municipality.', style='List Number')
document.add_paragraph('3. Guidance & Education Agent: Combines the object data and local rules to generate a simple, actionable instruction for the user.', style='List Number')
document.add_paragraph('4. Gamification Agent: Tracks user consistency, awards points for learning proper habits, and triggers reminders.', style='List Number')


document.add_heading('5. Target Audience & Customer Segments', level=2)
document.add_paragraph('Primary Users: Eco-conscious households and young adults.', style='List Bullet')
document.add_paragraph('B2B/Institutional Customers: Local governments and municipal waste management organizations.', style='List Bullet')

document.add_heading('6. Expected Impact & Metrics', level=2)
document.add_paragraph('Reduce Recycling Contamination: Lower the percentage of contaminated recycling batches in partner municipalities by 20% in the first year.', style='List Bullet')
document.add_paragraph('Increase Awareness: Educate thousands of households daily on sustainable consumption habits.', style='List Bullet')

document.save('Concept_Note.docx')
print("Successfully saved Concept_Note.docx")
